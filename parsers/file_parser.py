"""Lectura de texto desde TXT, MD, PDF, DOCX y XLSX (imports perezosos)."""
from pathlib import Path

from i18n import t

EXTENSIONES = ["txt", "md", "pdf", "docx", "xlsx"]


def extract_text(file_path: str) -> str:
    ext = Path(file_path).suffix.lower().lstrip(".")
    if ext in ("txt", "md"):
        return Path(file_path).read_text(encoding="utf-8", errors="ignore")
    if ext == "pdf":
        from pypdf import PdfReader
        from parsers.pdf_plan import a_markdown     # si parece un plan, recupera módulos y temas que el PDF aplana
        return a_markdown("\n".join((p.extract_text() or "") for p in PdfReader(file_path).pages))
    if ext == "docx":
        try:
            from docx import Document
        except ImportError:
            raise ValueError(t("DOCX no está disponible en esta versión. Usa MD, TXT, PDF o XLSX."))
        doc = Document(file_path)
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text for c in fila.cells))
        return "\n".join(partes)
    if ext == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        filas = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                filas.append(" ".join(str(c) for c in row if c is not None))
        return "\n".join(filas)
    raise ValueError(t("Formato no soportado: .{ext}", ext=ext))

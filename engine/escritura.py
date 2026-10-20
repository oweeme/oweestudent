"""Escritor de monografías, ensayos, tesis y libros.

Un escrito es un conjunto de secciones con meta de palabras. Se escribe con citas del tipo `[@clave]` (o `[@clave, p. 12]`,
o `[@uno; @otro]`) que apuntan a las lecturas registradas en «Lecturas»; la app las convierte al estilo elegido
(autor-año tipo APA, o numérico) y arma la lista de referencias con lo que realmente citaste. Se exporta a Markdown, HTML
(para imprimir a PDF) y Word.

La IA (opcional) puede sugerir fuentes, redactar un BORRADOR a partir de tus fichas y revisar la redacción, pero solo cita
lo que tú ya registraste: no inventa referencias.
"""
import html as _html
import re
from datetime import date

from engine import biblio
from i18n import marcar, t, tn

TIPOS = ("monografia", "ensayo", "tesis", "libro")
NOMBRES_TIPO = {"monografia": marcar("Monografía"), "ensayo": marcar("Ensayo"), "tesis": marcar("Tesis"), "libro": marcar("Libro")}
ESTILOS = ("apa", "numerico")
NOMBRES_ESTILO = {"apa": marcar("Autor y año (tipo APA)"), "numerico": marcar("Numérico [1]")}
METAS_POR_TIPO = {"monografia": 3000, "ensayo": 1500, "tesis": 20000, "libro": 40000}
_CITA = re.compile(r"\[(@[^\]]+)\]")
_UNA = re.compile(r"@\s*([\w:.\-]+)(?:\s*,\s*p+\.?\s*([\w\-–]+))?")


# ---------- plantillas ----------
def esqueleto(tipo: str, capitulos: int = 8) -> list[tuple[str, int, str]]:
    """(título, peso en %, guía de qué escribir) de cada sección, ya en el idioma de la interfaz."""
    if tipo == "monografia":
        return [(t("Introducción"), 10, t("Presenta el tema, por qué es importante y qué objetivo tiene tu trabajo.")),
                (t("Antecedentes"), 15, t("Qué se ha dicho antes sobre el tema (cita tus fuentes).")),
                (t("Desarrollo del tema"), 40, t("Explica el tema por partes, con ejemplos y citas.")),
                (t("Análisis"), 25, t("Tu propia reflexión: compara, relaciona y evalúa lo que expusiste.")),
                (t("Conclusiones"), 10, t("Responde al objetivo con lo más importante; no añadas ideas nuevas."))]
    if tipo == "ensayo":
        return [(t("Introducción y tesis"), 15, t("Plantea el problema y di claramente qué defiendes.")),
                (t("Primer argumento"), 20, t("Una idea que respalda tu tesis, con evidencia.")),
                (t("Segundo argumento"), 20, t("Otra idea que respalda tu tesis, con evidencia.")),
                (t("Tercer argumento"), 15, t("Otra idea, o un ejemplo que refuerce las anteriores.")),
                (t("Contraargumento"), 15, t("La mejor objeción a tu tesis y por qué no la derrumba.")),
                (t("Conclusión"), 15, t("Resume lo defendido y deja una idea final."))]
    if tipo == "tesis":
        return [(t("Resumen"), 3, t("Una página: problema, método, resultados principales y conclusión.")),
                (t("Introducción"), 10, t("Planteamiento del problema, objetivos, preguntas y justificación.")),
                (t("Marco teórico"), 25, t("Teorías y estudios previos que sostienen tu trabajo (cita tus lecturas).")),
                (t("Metodología"), 12, t("Qué hiciste, con quién y cómo, para que otra persona pueda repetirlo.")),
                (t("Resultados"), 20, t("Lo que encontraste, sin interpretarlo todavía.")),
                (t("Discusión"), 18, t("Qué significan los resultados y cómo se relacionan con la teoría.")),
                (t("Conclusiones y recomendaciones"), 12, t("Responde a tus preguntas; límites del estudio y futuras líneas."))]
    if tipo == "libro":
        n = max(1, capitulos)
        cap = [(t("Capítulo {n}", n=i + 1), round(90 / n), t("Idea central del capítulo, ejemplos y cierre que conduzca al siguiente.")) for i in range(n)]
        return [(t("Prólogo"), 5, t("Por qué escribes este libro y a quién se dirige."))] + cap + \
               [(t("Epílogo"), 5, t("Cierre: qué se llevará el lector."))]
    raise ValueError("tipo de escrito desconocido")


# ---------- escritos y secciones ----------
def crear_escrito(conn, plan_id, titulo, tipo="monografia", meta_palabras=None, fecha_limite=None, autor="", estilo="apa", capitulos=8) -> int:
    if tipo not in TIPOS or estilo not in ESTILOS:
        raise ValueError("tipo o estilo de escrito desconocido")
    meta = meta_palabras or METAS_POR_TIPO[tipo]
    limite = fecha_limite.isoformat() if hasattr(fecha_limite, "isoformat") else (date.fromisoformat(fecha_limite).isoformat() if fecha_limite else None)
    cur = conn.execute("INSERT INTO escritos(plan_id,titulo,tipo,estilo,autor,meta_palabras,fecha_limite) VALUES (?,?,?,?,?,?,?)",
                       (plan_id, titulo.strip(), tipo, estilo, autor.strip(), meta, limite))
    eid = cur.lastrowid
    for i, (tit, peso, guia) in enumerate(esqueleto(tipo, capitulos)):
        conn.execute("INSERT INTO secciones(escrito_id,orden,titulo,meta_palabras,notas) VALUES (?,?,?,?,?)",
                     (eid, i, tit, max(10, round(meta * peso / 100, -1)), guia))
    conn.commit()
    return eid


def escritos(conn, plan_id):
    return conn.execute("SELECT * FROM escritos WHERE plan_id=? ORDER BY id", (plan_id,)).fetchall()


def escrito(conn, escrito_id):
    return conn.execute("SELECT * FROM escritos WHERE id=?", (escrito_id,)).fetchone()


def editar_escrito(conn, escrito_id, **campos):
    ok = {k: v for k, v in campos.items() if k in ("titulo", "tipo", "estilo", "autor", "meta_palabras", "fecha_limite")}
    if "estilo" in ok and ok["estilo"] not in ESTILOS:
        raise ValueError("estilo de cita desconocido")
    if ok.get("fecha_limite"):
        ok["fecha_limite"] = ok["fecha_limite"].isoformat() if hasattr(ok["fecha_limite"], "isoformat") else date.fromisoformat(ok["fecha_limite"]).isoformat()
    if ok:
        conn.execute(f"UPDATE escritos SET {','.join(k + '=?' for k in ok)} WHERE id=?", list(ok.values()) + [escrito_id])
        conn.commit()


def borrar_escrito(conn, escrito_id):
    conn.execute("DELETE FROM escritos WHERE id=?", (escrito_id,))
    conn.commit()


def secciones(conn, escrito_id):
    return conn.execute("SELECT * FROM secciones WHERE escrito_id=? ORDER BY orden, id", (escrito_id,)).fetchall()


def seccion(conn, seccion_id):
    return conn.execute("SELECT * FROM secciones WHERE id=?", (seccion_id,)).fetchone()


def agregar_seccion(conn, escrito_id, titulo, meta_palabras=0, notas="") -> int:
    orden = conn.execute("SELECT COALESCE(MAX(orden)+1,0) FROM secciones WHERE escrito_id=?", (escrito_id,)).fetchone()[0]
    cur = conn.execute("INSERT INTO secciones(escrito_id,orden,titulo,meta_palabras,notas) VALUES (?,?,?,?,?)",
                       (escrito_id, orden, titulo.strip(), meta_palabras, notas))
    conn.commit()
    return cur.lastrowid


def editar_seccion(conn, seccion_id, **campos):
    ok = {k: v for k, v in campos.items() if k in ("titulo", "texto", "meta_palabras", "notas")}
    if ok:
        conn.execute(f"UPDATE secciones SET {','.join(k + '=?' for k in ok)} WHERE id=?", list(ok.values()) + [seccion_id])
        conn.commit()


def mover_seccion(conn, seccion_id, delta: int):
    """Sube (-1) o baja (+1) una sección intercambiándola con su vecina."""
    s = seccion(conn, seccion_id)
    lista = [x["id"] for x in secciones(conn, s["escrito_id"])]
    i = lista.index(seccion_id)
    j = i + delta
    if 0 <= j < len(lista):
        lista[i], lista[j] = lista[j], lista[i]
        for orden, sid in enumerate(lista):
            conn.execute("UPDATE secciones SET orden=? WHERE id=?", (orden, sid))
        conn.commit()


def borrar_seccion(conn, seccion_id):
    conn.execute("DELETE FROM secciones WHERE id=?", (seccion_id,))
    conn.commit()


# ---------- palabras y avance ----------
def contar_palabras(texto: str) -> int:
    """Palabras del texto sin contar las marcas de cita ni los símbolos de formato."""
    limpio = _CITA.sub("", texto or "")
    cjk = len(re.findall(r"[぀-ヿ㐀-䶿一-鿿]", limpio))      # chino/japonés: un carácter ≈ una palabra
    return len(re.findall(r"[^\W_]+(?:['’\-][^\W_]+)*", re.sub(r"[぀-ヿ㐀-䶿一-鿿]", " ", limpio))) + cjk


def progreso(conn, escrito_id) -> dict:
    e = escrito(conn, escrito_id)
    filas = [(s, contar_palabras(s["texto"])) for s in secciones(conn, escrito_id)]
    total = sum(n for _, n in filas)
    meta = e["meta_palabras"] or sum(s["meta_palabras"] or 0 for s, _ in filas)
    dias = None
    if e["fecha_limite"]:
        dias = (date.fromisoformat(e["fecha_limite"]) - date.today()).days
    faltan = max(0, meta - total)
    return {"palabras": total, "meta": meta, "porcentaje": min(100, round(100 * total / meta)) if meta else 0, "dias": dias,
            "palabras_por_dia": (faltan / dias if dias and dias > 0 else None),
            "secciones": [{"id": s["id"], "titulo": s["titulo"], "palabras": n, "meta": s["meta_palabras"] or 0} for s, n in filas]}


# ---------- citas y referencias ----------
def claves_del_plan(conn, plan_id) -> dict:
    """{clave: lectura} de las lecturas del plan. Usa la clave BibTeX si la hay; si no, se genera (estable y única)."""
    usadas: set[str] = set()
    out = {}
    for r in conn.execute("SELECT * FROM lecturas WHERE plan_id=? ORDER BY id", (plan_id,)):
        d = dict(r)
        c = d.get("clave_bib") if d.get("clave_bib") and d["clave_bib"] not in usadas else biblio.clave_para(d, usadas)
        usadas.add(c)
        out[c] = r
    return out


def citas_en(texto: str) -> list[tuple[str, str | None]]:
    """[(clave, página)] de todas las citas del texto, en orden de aparición."""
    out = []
    for m in _CITA.finditer(texto or ""):
        for parte in m.group(1).split(";"):
            u = _UNA.match(parte.strip())
            if u:
                out.append((u.group(1), u.group(2)))
    return out


def _apellidos(lectura) -> list[str]:
    return [a.split(",")[0].strip() for a in (lectura["autores"] or "").split(";") if a.strip()]


def _lista(partes: list[str]) -> str:
    if len(partes) <= 1:
        return "".join(partes)
    return t("{lista} y {ultimo}", lista=", ".join(partes[:-1]), ultimo=partes[-1])


def cita_corta(lectura, pagina=None) -> str:
    """Parte de la cita en el texto estilo autor-año: «Ruiz y Quispe, 2021, p. 4»."""
    ap = _apellidos(lectura)
    quien = (ap[0] if len(ap) == 1 else t("{a} y {b}", a=ap[0], b=ap[1]) if len(ap) == 2 else t("{a} et al.", a=ap[0])) if ap \
        else f"“{(lectura['titulo'] or '')[:30].strip()}”"
    anio = str(lectura["anio"]) if lectura["anio"] else t("s. f.")
    pag = f", {t('p. {n}', n=pagina)}" if pagina else ""
    return f"{quien}, {anio}{pag}"


def referencia(lectura, estilo="apa", numero=None) -> str:
    """Una entrada de la lista de referencias, con *cursiva* para la fuente."""
    ap = []
    for a in (lectura["autores"] or "").split(";"):
        a = a.strip()
        if not a:
            continue
        ape, _, nom = a.partition(",")
        iniciales = " ".join(f"{p[0]}." for p in re.split(r"[\s\-]+", nom.strip()) if p)
        ap.append(f"{ape.strip()}, {iniciales}".strip(", "))
    quien = _lista(ap)
    anio = f"({lectura['anio']})" if lectura["anio"] else f"({t('s. f.')})"
    enlace = (f"https://doi.org/{lectura['doi']}" if lectura["doi"] and not lectura["doi"].startswith("http") else (lectura["doi"] or lectura["url"] or ""))
    fuente = f" *{lectura['fuente']}*." if lectura["fuente"] else ""
    cuerpo = f"{quien + ' ' if quien else ''}{anio}. {lectura['titulo'].rstrip('.')}.{fuente}{' ' + enlace if enlace else ''}"
    return f"[{numero}] {cuerpo}" if estilo == "numerico" and numero else cuerpo


def _convertir_citas(texto: str, claves: dict, estilo: str, orden: list[str], sin_fuente: set[str]) -> str:
    def una(m):
        partes, salida = m.group(1).split(";"), []
        for p in partes:
            u = _UNA.match(p.strip())
            if not u:
                continue
            clave, pag = u.group(1), u.group(2)
            lect = claves.get(clave)
            if lect is None:
                sin_fuente.add(clave)
                salida.append(f"{clave}?")
                continue
            if clave not in orden:
                orden.append(clave)
            salida.append(cita_corta(lect, pag) if estilo == "apa" else str(orden.index(clave) + 1) + (f", {t('p. {n}', n=pag)}" if pag else ""))
        if not salida:
            return m.group(0)
        return f"({'; '.join(salida)})" if estilo == "apa" else f"[{'; '.join(salida)}]"
    return _CITA.sub(una, texto or "")


def compilar(conn, escrito_id) -> dict:
    """El documento final: {titulo, autor, tipo, bloques:[{titulo, parrafos}], referencias:[…], sin_fuente:[claves], palabras}."""
    e = escrito(conn, escrito_id)
    claves = claves_del_plan(conn, e["plan_id"])
    orden: list[str] = []
    sin_fuente: set[str] = set()
    bloques, palabras = [], 0
    for s in secciones(conn, escrito_id):
        texto = _convertir_citas(s["texto"] or "", claves, e["estilo"], orden, sin_fuente)
        palabras += contar_palabras(s["texto"])
        parrafos = [p.strip() for p in re.split(r"\n\s*\n", texto) if p.strip()]
        bloques.append({"titulo": s["titulo"], "parrafos": parrafos})
    usadas = [claves[c] for c in orden]
    if e["estilo"] == "apa":
        refs = sorted((referencia(l) for l in usadas), key=lambda r: r.lower())
    else:
        refs = [referencia(l, "numerico", i + 1) for i, l in enumerate(usadas)]
    return {"titulo": e["titulo"], "autor": e["autor"] or "", "tipo": e["tipo"], "bloques": bloques, "referencias": refs,
            "sin_fuente": sorted(sin_fuente), "palabras": palabras}


# ---------- exportar ----------
def a_markdown(doc: dict) -> str:
    partes = [f"# {doc['titulo']}"]
    if doc["autor"]:
        partes.append(f"*{doc['autor']}*")
    for b in doc["bloques"]:
        partes.append(f"## {b['titulo']}")
        partes += b["parrafos"]
    if doc["referencias"]:
        partes.append(f"## {t('Referencias')}")
        partes += doc["referencias"]
    return "\n\n".join(partes) + "\n"


def _inline(s: str) -> str:
    s = _html.escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
    return re.sub(r"\*(.+?)\*", r"<em>\1</em>", s)


def a_html(doc: dict) -> str:
    cuerpo = [f"<h1>{_html.escape(doc['titulo'])}</h1>"]
    if doc["autor"]:
        cuerpo.append(f'<p class="autor">{_html.escape(doc["autor"])}</p>')
    for b in doc["bloques"]:
        cuerpo.append(f"<h2>{_html.escape(b['titulo'])}</h2>")
        cuerpo += [f"<p>{_inline(p).replace(chr(10), '<br>')}</p>" for p in b["parrafos"]]
    if doc["referencias"]:
        cuerpo.append(f"<h2>{_html.escape(t('Referencias'))}</h2>")
        cuerpo += [f'<p class="ref">{_inline(r)}</p>' for r in doc["referencias"]]
    css = ("body{font-family:Georgia,'Times New Roman',serif;max-width:42em;margin:2em auto;padding:0 1em;line-height:1.6;color:#111}"
           "h1{text-align:center}.autor{text-align:center;color:#555}p{text-align:justify}.ref{padding-left:2em;text-indent:-2em;text-align:left}"
           "@media print{body{margin:0;max-width:none}}")
    return f'<!doctype html>\n<html lang="es"><head><meta charset="utf-8"><title>{_html.escape(doc["titulo"])}</title><style>{css}</style></head>\n<body>\n' + \
           "\n".join(cuerpo) + "\n</body></html>\n"


def docx_disponible() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


def a_docx(doc: dict, ruta) -> None:
    """Documento de Word. Necesita python-docx (no está en Android)."""
    try:
        from docx import Document
        from docx.shared import Cm, Pt
    except ImportError as ex:
        raise RuntimeError(t("Word no está disponible en este dispositivo. Usa Markdown o HTML.")) from ex
    d = Document()
    d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Times New Roman", Pt(12)
    d.add_heading(doc["titulo"], 0)
    if doc["autor"]:
        d.add_paragraph(doc["autor"])

    def con_formato(par, texto):
        for trozo in re.split(r"(\*\*.+?\*\*|\*.+?\*)", texto):
            if trozo.startswith("**") and trozo.endswith("**") and len(trozo) > 4:
                par.add_run(trozo[2:-2]).bold = True
            elif trozo.startswith("*") and trozo.endswith("*") and len(trozo) > 2:
                par.add_run(trozo[1:-1]).italic = True
            elif trozo:
                par.add_run(trozo)
    for b in doc["bloques"]:
        d.add_heading(b["titulo"], 1)
        for p in b["parrafos"]:
            con_formato(d.add_paragraph(), p)
    if doc["referencias"]:
        d.add_heading(t("Referencias"), 1)
        for r in doc["referencias"]:
            par = d.add_paragraph()
            par.paragraph_format.left_indent, par.paragraph_format.first_line_indent = Cm(1.25), Cm(-1.25)
            con_formato(par, r)
    d.save(str(ruta))


# ---------- revisión de la redacción (sin IA) ----------
def revisar_seccion(texto: str, meta: int = 0, claves: set[str] | None = None) -> list[str]:
    """Avisos prácticos (ya traducidos) sobre una sección. Sin IA: largo de frases y párrafos, citas y meta de palabras."""
    avisos = []
    claves = claves or set()
    n = contar_palabras(texto)
    if n == 0:
        return [t("La sección está vacía.")]
    limpio = _CITA.sub("", texto)
    largas = [f for f in re.split(r"(?<=[.!?。！？])\s+", limpio) if len(f.split()) > 40]
    if largas:
        avisos.append(tn(len(largas), "Hay {n} frase de más de 40 palabras: divídela.", "Hay {n} frases de más de 40 palabras: divídelas."))
    parrafos = [p for p in re.split(r"\n\s*\n", texto) if p.strip()]
    gordos = [p for p in parrafos if contar_palabras(p) > 220]
    if gordos:
        avisos.append(tn(len(gordos), "Hay {n} párrafo de más de 220 palabras: sepáralo en ideas.", "Hay {n} párrafos de más de 220 palabras: sepáralos en ideas."))
    sin_cita = [p for p in parrafos if contar_palabras(p) > 100 and not _CITA.search(p)]
    if sin_cita:
        avisos.append(tn(len(sin_cita), "{n} párrafo largo no cita ninguna fuente: ¿de dónde sale lo que afirmas?",
                         "{n} párrafos largos no citan ninguna fuente: ¿de dónde sale lo que afirmas?"))
    desconocidas = sorted({c for c, _ in citas_en(texto) if claves and c not in claves})
    if desconocidas:
        avisos.append(t("Citas sin lectura registrada: {lista}. Añádelas en «Lecturas» o corrige la clave.", lista=", ".join(desconocidas)))
    palabras = [w.lower() for w in re.findall(r"[^\W\d_]{6,}", limpio)]
    if palabras:
        repetida = max(set(palabras), key=palabras.count)
        if palabras.count(repetida) >= max(5, n // 40):
            avisos.append(t("Repites mucho «{w}» ({n} veces): busca sinónimos o reformula.", w=repetida, n=palabras.count(repetida)))
    if meta:
        if n < meta * 0.5:
            avisos.append(t("Llevas {n} de {meta} palabras de la meta de esta sección.", n=n, meta=meta))
        elif n > meta * 1.5:
            avisos.append(t("Pasas de la meta: {n} palabras frente a {meta}. ¿Se puede condensar?", n=n, meta=meta))
    return avisos or [t("Se ve ordenada. Léela en voz alta para detectar lo que suene raro.")]


# ---------- fuentes y borrador con IA ----------
def fuentes_para(conn, indice, plan_id: int, consulta: str, k: int = 5) -> list[dict]:
    """Pasajes de tus documentos, notas y fichas relacionados con `consulta`, y la clave para citarlos si están en «Lecturas»."""
    from engine import busqueda as bq
    claves = claves_del_plan(conn, plan_id)
    por_id = {r["id"]: c for c, r in claves.items()}
    por_material = {r["material_id"]: c for c, r in claves.items() if r["material_id"]}
    out = []
    for f in bq.buscar(conn, indice, [plan_id], consulta, k):
        clave = por_id.get(f.ref_id) if f.origen == "lectura" else por_material.get(f.ref_id)
        out.append({"fuente": f, "clave": clave})
    return out


def prompt_borrador(seccion_titulo: str, puntos: str, fuentes: list[dict], por_fragmento: int = 600) -> str:
    frag = "\n\n".join(f"[{i}] ({x['fuente'].etiqueta()})\n{x['fuente'].texto[:por_fragmento]}" for i, x in enumerate(fuentes, start=1))
    return (f"FRAGMENTOS:\n{frag}\n\nSECCIÓN: {seccion_titulo}\nIDEAS QUE EL AUTOR QUIERE DESARROLLAR: {puntos or '(ninguna: resume lo esencial)'}\n\n"
            "Escribe 1 o 2 párrafos académicos y claros para esa sección. Usa SOLO la información de los FRAGMENTOS y termina cada "
            "afirmación con la cita del fragmento, como [1]. No inventes datos ni fuentes. Si los fragmentos no alcanzan, escribe solo lo que sí respaldan.")


_SIS_BORRADOR = ("Eres un asistente de escritura académica. Redactas borradores sobrios y claros en el idioma del autor, "
                 "apoyados únicamente en los fragmentos dados y citándolos con [n]. No usas conocimiento propio.")


def borrador_ia(conn, indice, plan_id: int, seccion_titulo: str, puntos: str, llm=None) -> dict:
    """Borrador con citas `[@clave]` para las fuentes registradas. Devuelve {texto, sin_registrar:[etiquetas], sin_fuentes}."""
    from engine import busqueda as bq
    fuentes = fuentes_para(conn, indice, plan_id, f"{seccion_titulo} {puntos}".strip(), 4)
    if not fuentes:
        return {"texto": "", "sin_registrar": [], "sin_fuentes": True, "no_alcanza": False}
    if llm is None:
        from engine import ai
        llm = ai._llm
    crudo = llm(prompt_borrador(seccion_titulo, puntos, fuentes), _SIS_BORRADOR, 450)
    rev = bq.revisar_respuesta(crudo, len(fuentes))
    if not rev["encontrada"]:
        return {"texto": "", "sin_registrar": [], "sin_fuentes": False, "no_alcanza": True}
    texto, no_citables = rev["texto"], []

    def sustituir(m):
        i = int(m.group(1))
        clave = fuentes[i - 1]["clave"]
        if clave:
            return f" [@{clave}]"
        et = fuentes[i - 1]["fuente"].etiqueta()
        if et not in no_citables:
            no_citables.append(et)
        return ""
    texto = re.sub(r"[ ]?\[(\d+)\]", sustituir, texto)
    return {"texto": texto.strip(), "sin_registrar": no_citables, "sin_fuentes": False, "no_alcanza": False}

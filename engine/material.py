"""Material de estudio: adjuntar sin duplicar (hash), leer por páginas y convertir tablas en tarjetas."""
import csv
import hashlib
import shutil
from pathlib import Path

from database.db import DATOS

TIPOS = {"pdf": "pdf", "docx": "docx", "xlsx": "xlsx", "csv": "csv", "txt": "txt", "md": "md"}
CARPETA = DATOS / "materiales"
_cache: dict = {}


def tipo_de(ruta: str) -> str:
    return TIPOS.get(Path(ruta).suffix.lower().lstrip("."), "otro")


def hash_archivo(ruta: str) -> str:
    h = hashlib.sha256()
    with open(ruta, "rb") as f:
        for bloque in iter(lambda: f.read(1 << 20), b""):
            h.update(bloque)
    return h.hexdigest()


def copiar_a_app(ruta: str, hash_: str) -> str:
    """Guarda UNA copia por contenido (nombre = hash): el mismo libro nunca se duplica."""
    CARPETA.mkdir(parents=True, exist_ok=True)
    destino = CARPETA / f"{hash_}{Path(ruta).suffix.lower()}"
    if not destino.exists():
        shutil.copy2(ruta, destino)
    return str(destino)


# ---------------- lectura por páginas ----------------
def _trozos(texto: str, tam: int = 2500) -> list[str]:
    paginas, actual = [], ""
    for parrafo in texto.split("\n"):
        if len(actual) + len(parrafo) > tam and actual:
            paginas.append(actual.strip())
            actual = ""
        actual += parrafo + "\n"
    if actual.strip():
        paginas.append(actual.strip())
    return paginas or [""]


def _paginas_cache(ruta: str):
    clave = (ruta, Path(ruta).stat().st_mtime)
    if clave not in _cache:
        _cache.clear()  # una sola entrada: evita acumular memoria
        t = tipo_de(ruta)
        if t == "pdf":
            from pypdf import PdfReader
            _cache[clave] = ("pdf", PdfReader(ruta))
        elif t == "docx":
            from parsers.file_parser import extract_text
            _cache[clave] = ("lista", _trozos(extract_text(ruta)))
        elif t in ("xlsx", "csv"):
            hojas, _ = leer_tabla(ruta)
            paginas = []
            for i in range(len(hojas)):
                _, filas = leer_tabla(ruta, i)
                for j in range(0, max(1, len(filas)), 40):
                    paginas.append(f"[{hojas[i]}]\n" + "\n".join(" | ".join(f) for f in filas[j:j + 40]))
            _cache[clave] = ("lista", paginas or [""])
        else:
            _cache[clave] = ("lista", _trozos(Path(ruta).read_text(encoding="utf-8", errors="ignore")))
    return _cache[clave]


def num_paginas(ruta: str) -> int:
    tipo, datos = _paginas_cache(ruta)
    return len(datos.pages) if tipo == "pdf" else len(datos)


def leer_pagina(ruta: str, n: int) -> str:
    """Texto de la página `n` (empieza en 1)."""
    tipo, datos = _paginas_cache(ruta)
    total = len(datos.pages) if tipo == "pdf" else len(datos)
    n = max(1, min(n, total))
    if tipo == "pdf":
        return (datos.pages[n - 1].extract_text() or "").strip() or "(Esta página no tiene texto seleccionable: ¿es un escaneo?)"
    return datos[n - 1]


# ---------------- tablas -> tarjetas (vocabulario) ----------------
def leer_tabla(ruta: str, hoja: int = 0) -> tuple[list[str], list[list[str]]]:
    """(nombres de hojas, filas de la hoja elegida). CSV cuenta como una sola hoja."""
    if tipo_de(ruta) == "csv":
        with open(ruta, newline="", encoding="utf-8-sig", errors="ignore") as f:
            filas = [[c.strip() for c in fila] for fila in csv.reader(f)]
        return ["CSV"], [f for f in filas if any(f)]
    import openpyxl
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        nombres = wb.sheetnames
        ws = wb[nombres[min(hoja, len(nombres) - 1)]]
        filas = [["" if c is None else str(c).strip() for c in fila] for fila in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return nombres, [f for f in filas if any(f)]


def a_tarjetas(filas: list[list[str]], frente: int, reverso: int, extra: int | None = None,
               con_cabecera: bool = True, ambos_sentidos: bool = False) -> list[tuple[str, str]]:
    """Convierte filas en (pregunta, respuesta). `extra` (p. ej. pinyin) se añade a la respuesta."""
    pares = []
    for fila in filas[1 if con_cabecera else 0:]:
        g = lambda i: fila[i].strip() if i is not None and i < len(fila) else ""
        f, r, e = g(frente), g(reverso), g(extra)
        if not f or not r:
            continue
        pares.append((f, f"{r} ({e})" if e else r))
        if ambos_sentidos:
            pares.append((r, f"{f} ({e})" if e else f))
    return pares

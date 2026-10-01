import shutil
import time

import pytest

from conftest import EJEMPLO
from database import repo
from database.db import init_db
from engine import material
from engine.merge import fusionar


@pytest.fixture
def bd():
    return init_db(":memory:")


@pytest.fixture(autouse=True)
def carpeta_material(tmp_path, monkeypatch):
    monkeypatch.setattr(material, "CARPETA", tmp_path / "materiales")


# ---------- carpetas y asignaturas ----------
def test_carpetas_y_asignaturas(bd):
    c = repo.crear_carpeta(bd, "3º Secundaria")
    mat = repo.crear_plan(bd, "Matemática", 1, carpeta_id=c)
    cie = repo.crear_plan(bd, "Ciencias", 1, carpeta_id=c)
    suelta = repo.crear_plan(bd, "Inglés", 1)
    assert [p["titulo"] for p in repo.planes_de_carpeta(bd, 1, c)] == ["Matemática", "Ciencias"]
    assert [p["titulo"] for p in repo.planes_de_carpeta(bd, 1, None)] == ["Inglés"]
    repo.mover_plan(bd, suelta, c)
    assert len(repo.planes_de_carpeta(bd, 1, c)) == 3
    repo.borrar_carpeta(bd, c)
    assert repo.carpetas(bd, 1) == [] and repo.planes(bd, 1) == []


def test_importar_dos_veces_no_duplica(bd):
    pid = repo.importar_archivo(bd, str(EJEMPLO))
    n_u, n_c = len(repo.unidades(bd, pid)), len(repo.cronograma(bd, pid))
    repo.importar_archivo(bd, str(EJEMPLO), plan_id=pid)
    assert len(repo.unidades(bd, pid)) == n_u and len(repo.cronograma(bd, pid)) == n_c


def test_importar_en_carpeta(bd):
    c = repo.crear_carpeta(bd, "Máster")
    pid = repo.importar_archivo(bd, str(EJEMPLO), carpeta_id=c)
    assert bd.execute("select carpeta_id from planes where id=?", (pid,)).fetchone()[0] == c


def test_migracion_quita_cronograma_duplicado(tmp_path):
    ruta = str(tmp_path / "v.db")
    c = init_db(ruta)
    pid = repo.importar_archivo(c, str(EJEMPLO))
    for _ in range(2):  # duplicados como los de una versión anterior
        c.execute("INSERT INTO cronograma(plan_id,dia,bloque,actividad) VALUES (?,?,?,?)", (pid, "Lunes", "B", "Dup"))
    c.commit(); c.close()
    c2 = init_db(ruta)
    assert c2.execute("select count(*) from cronograma where actividad='Dup'").fetchone()[0] == 1


# ---------- material ----------
def test_material_sin_duplicados(bd, tmp_path):
    pid = repo.crear_plan(bd, "Redes", 1)
    f1 = tmp_path / "libro.txt"; f1.write_text("contenido del libro " * 50)
    f2 = tmp_path / "copia_del_libro.txt"; shutil.copy(f1, f2)          # mismo contenido, otro nombre
    h1, h2 = material.hash_archivo(str(f1)), material.hash_archivo(str(f2))
    assert h1 == h2
    r1, r2 = material.copiar_a_app(str(f1), h1), material.copiar_a_app(str(f2), h2)
    assert r1 == r2                                                      # una sola copia física
    mid = repo.agregar_material(bd, pid, "Libro", "txt", h1, "libro.txt", r1, f1.stat().st_size)
    assert repo.material_por_hash(bd, h2, pid)["id"] == mid              # se detecta el duplicado
    assert len(list((tmp_path / "materiales").iterdir())) == 1


def test_leer_txt_y_paginas(tmp_path):
    f = tmp_path / "l.md"
    f.write_text("\n".join(f"Línea {i} " + "x" * 80 for i in range(120)))
    assert material.num_paginas(str(f)) >= 3
    assert "Línea 0" in material.leer_pagina(str(f), 1)
    assert material.leer_pagina(str(f), 999)  # fuera de rango: última página


def test_leer_docx(tmp_path):
    import docx
    d = docx.Document(); d.add_paragraph("Hola desde Word"); d.save(tmp_path / "a.docx")
    assert "Hola desde Word" in material.leer_pagina(str(tmp_path / "a.docx"), 1)


def _pdf_minimo(texto: str) -> bytes:
    """PDF válido de una página con `texto` (con tabla xref calculada)."""
    flujo = f"BT /F1 12 Tf 20 100 Td ({texto}) Tj ET".encode()
    objs = [b"<</Type/Catalog/Pages 2 0 R>>", b"<</Type/Pages/Kids[3 0 R]/Count 1>>",
            b"<</Type/Page/Parent 2 0 R/MediaBox[0 0 300 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
            b"<</Length %d>>\nstream\n" % len(flujo) + flujo + b"\nendstream",
            b"<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offs)
    return out + b"trailer<</Size %d/Root 1 0 R>>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)


def test_leer_pdf(tmp_path):
    f = tmp_path / "a.pdf"; f.write_bytes(_pdf_minimo("Texto del PDF"))
    assert material.num_paginas(str(f)) == 1 and "Texto del PDF" in material.leer_pagina(str(f), 1)


def test_excel_a_tarjetas_vocabulario(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "HSK1"
    ws.append(["Chino", "Español", "Pinyin"]); ws.append(["你好", "hola", "nǐ hǎo"]); ws.append(["谢谢", "gracias", "xièxie"])
    ws.append(["", "vacío", ""])
    wb.create_sheet("Alemán").append(["Wort", "Bedeutung"])
    wb.save(tmp_path / "v.xlsx")
    hojas, filas = material.leer_tabla(str(tmp_path / "v.xlsx"), 0)
    assert hojas == ["HSK1", "Alemán"] and len(filas) == 4
    assert material.a_tarjetas(filas, 0, 1, 2) == [("你好", "hola (nǐ hǎo)"), ("谢谢", "gracias (xièxie)")]
    assert len(material.a_tarjetas(filas, 0, 1, 2, ambos_sentidos=True)) == 4


def test_csv_a_tarjetas(tmp_path):
    f = tmp_path / "v.csv"; f.write_text("Wort,Bedeutung\nHaus,casa\nBaum,árbol\n", encoding="utf-8")
    _, filas = material.leer_tabla(str(f))
    assert material.a_tarjetas(filas, 0, 1) == [("Haus", "casa"), ("Baum", "árbol")]


def test_tema_para_tarjetas_reutiliza(bd):
    pid = repo.crear_plan(bd, "Alemán", 1)
    a, b = repo.tema_para_tarjetas(bd, pid), repo.tema_para_tarjetas(bd, pid)
    assert a == b


# ---------- mezcla con tablas nuevas ----------
def test_mezcla_carpetas_material_y_clave(tmp_path):
    A = init_db(str(tmp_path / "A.db"))
    c = repo.crear_carpeta(A, "Máster"); pid = repo.crear_plan(A, "Management", 1, carpeta_id=c)
    mid = repo.agregar_material(A, pid, "Libro", "pdf", "h123", "l.pdf", "/ruta/A/l.pdf", 10)
    repo.agregar_nota_material(A, mid, 3, "conclusión A")
    shutil.copy(tmp_path / "A.db", tmp_path / "B.db"); B = init_db(str(tmp_path / "B.db"))
    time.sleep(0.02)
    repo.agregar_nota_material(B, mid, 5, "conclusión B"); repo.renombrar_carpeta(B, c, "Máster Global")
    repo.progreso_lectura(B, mid, 7, 100); repo.fijar_clave(A, 1, "claveA")
    fusionar(A, str(tmp_path / "B.db")); fusionar(B, str(tmp_path / "A.db"))
    for X in (A, B):
        assert X.execute("select nombre from carpetas").fetchone()[0] == "Máster Global"
        assert X.execute("select count(*) from notas_material").fetchone()[0] == 2
        assert X.execute("select pagina_actual from materiales").fetchone()[0] == 7
    # la clave de acceso NO viaja en la mezcla sobre un perfil existente
    assert repo.tiene_clave(A, 1) and not repo.tiene_clave(B, 1)


def test_merge_respeta_fk_opcional(tmp_path):
    A = init_db(str(tmp_path / "A.db")); B = init_db(str(tmp_path / "B.db"))
    pid = repo.crear_plan(A, "X", 1); u = repo.agregar_unidad(A, pid, "U")
    repo.agregar_material(A, pid, "M", "txt", "h", "m.txt", "/m.txt", 1, unidad_id=u)
    fusionar(B, str(tmp_path / "A.db"))
    assert B.execute("select count(*) from materiales").fetchone()[0] == 1
    assert B.execute("select unidad_id from materiales").fetchone()[0] is not None
